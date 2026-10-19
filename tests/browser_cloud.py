"""Static Pages E2E with real PostgreSQL RPCs behind a local test gateway.

HARBOR_TEST_DSN must name a dedicated harbor_test* DB; hosted Supabase is not mocked as ready.
"""
import functools
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
from contextlib import contextmanager
from http.server import SimpleHTTPRequestHandler,ThreadingHTTPServer,BaseHTTPRequestHandler
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import httpx
import psycopg
from psycopg.types.json import Jsonb
from playwright.sync_api import sync_playwright,expect
from openpyxl import load_workbook
from tests.fixtures import sample_rows,workbook

DSN=os.environ.get('HARBOR_TEST_DSN')
if not DSN:raise SystemExit('Set HARBOR_TEST_DSN to a dedicated harbor_test* PostgreSQL database')
OUTPUT=Path(os.environ.get('OUTPUT_DIR',ROOT/'test-results/free'))
OUTPUT.mkdir(parents=True,exist_ok=True)
created_workspaces=[]
RPC_ARGS={
 'harbor_create_workspace':{'p_token'},
 'harbor_snapshot':{'p_token','p_known_revision'},
 'harbor_commit':{'p_token','p_filename','p_file_hash','p_rows','p_accept_updates','p_skip_invalid','p_invalid','p_duplicates'},
 'harbor_history':{'p_token','p_batch_id'},
}


class QuietStatic(SimpleHTTPRequestHandler):
    def log_message(self,*args):pass


class RPCGateway(BaseHTTPRequestHandler):
    def log_message(self,*args):pass
    def do_POST(self):
        name=self.path.rsplit('/',1)[-1]
        try:
            args=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
            if name not in RPC_ARGS or not set(args).issubset(RPC_ARGS[name]):raise ValueError('Unknown RPC arguments')
            with psycopg.connect(DSN) as conn:
                conn.execute('set role anon')
                statement='select public.'+name+'('+','.join(k+' => %s' for k in args)+')'
                result=conn.execute(statement,[Jsonb(v) if isinstance(v,(dict,list)) else v for v in args.values()]).fetchone()[0]
            if name=='harbor_create_workspace':created_workspaces.append(hashlib.sha256(args['p_token'].encode()).hexdigest())
            self.send_response(200);body=json.dumps(result).encode()
        except psycopg.Error as error:
            self.send_response(400);body=json.dumps({'message':error.diag.message_primary,'code':error.sqlstate}).encode()
        except Exception as error:
            self.send_response(400);body=json.dumps({'message':str(error)}).encode()
        self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(body)));self.end_headers();self.wfile.write(body)


@contextmanager
def server(handler):
    http=ThreadingHTTPServer(('127.0.0.1',0),handler)
    thread=threading.Thread(target=http.serve_forever,daemon=True);thread.start()
    try:yield f'http://127.0.0.1:{http.server_port}'
    finally:http.shutdown();http.server_close();thread.join(timeout=5)


def run():
    with psycopg.connect(DSN) as conn:
        if not conn.info.dbname.startswith('harbor_test'):raise RuntimeError('Use a dedicated harbor_test* database')
        conn.execute((ROOT/'supabase/schema.sql').read_text())
    subprocess.run([sys.executable,'scripts/build_pages.py'],cwd=ROOT,check=True)
    with tempfile.TemporaryDirectory(prefix='harbor-pages-e2e-') as directory:
        # A project subpath catches root-relative assets that would fail on GitHub Pages.
        project=Path(directory)/'Test-for-EDI-Developer'
        shutil.copytree(ROOT/'dist',project)
        handler=functools.partial(QuietStatic,directory=directory)
        with server(handler) as static,server(RPCGateway) as gateway,sync_playwright() as p:
            options={'headless':True,'args':['--no-sandbox']}
            if os.environ.get('CHROMIUM_PATH'):options['executable_path']=os.environ['CHROMIUM_PATH']
            browser=p.chromium.launch(**options)
            errors=[]
            def wire(context):
                context.on('page',lambda page:page.on('pageerror',lambda error:errors.append(str(error))))
                def route_rpc(route):
                    endpoint=route.request.url.split('/rest/v1/rpc/')[1]
                    result=httpx.post(gateway+'/rest/v1/rpc/'+endpoint,content=route.request.post_data,headers={'Content-Type':'application/json'},trust_env=False,timeout=30)
                    route.fulfill(status=result.status_code,content_type='application/json',body=result.content)
                context.route('https://harbor-test.supabase.co/rest/v1/rpc/**',route_rpc)
            context=browser.new_context(viewport={'width':1440,'height':1000});wire(context)
            context.grant_permissions(['clipboard-read','clipboard-write'])
            page=context.new_page();base=static+'/Test-for-EDI-Developer/'
            page.goto(base,wait_until='networkidle')
            expect(page.locator('#setup-screen')).to_be_visible()
            page.locator('#copy-setup-sql').click()
            expect(page.locator('#sql-copy-status')).to_contain_text('คัดลอกแล้ว')
            assert 'create or replace function public.harbor_commit' in page.evaluate('async () => await navigator.clipboard.readText()')
            page.screenshot(path=OUTPUT/'free-setup.png',full_page=True)
            page.locator('#try-demo').click()
            expect(page.locator('#nav-count')).to_have_text('12')
            expect(page.locator('#share-workspace')).to_be_disabled()
            expect(page.locator('#local-demo-note')).to_be_visible()

            # Shared mode: one workspace, two independent browsers, real database.
            page.locator('#connection-settings').click()
            page.locator('#supabase-url').fill('https://harbor-test.supabase.co')
            page.locator('#supabase-key').fill('sb_publishable_test_key_for_browser_only')
            page.locator('#connect-submit').click()
            expect(page.locator('#connection-badge')).to_have_text('ข้อมูลร่วมกัน · Supabase Free')
            expect(page.locator('#nav-count')).to_have_text('0')
            shared_link=page.url
            assert '#join=' in shared_link
            page.locator('#share-workspace').click()
            assert page.evaluate('async () => await navigator.clipboard.readText()')==shared_link
            rows=sample_rows(100);rows.append({'port':'Invalid synthetic record'})
            def upload(page,rows,name='synthetic.xlsx'):
                page.locator('[data-import]').first.click()
                page.locator('#file-input').set_input_files({'name':name,'mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':workbook(rows)})
            upload(page,rows)
            expect(page.locator('#import-preview')).to_be_visible(timeout=15000)
            expect(page.locator('#confirm-import')).to_be_disabled()
            page.locator('#accept-invalid').check();page.locator('#confirm-import').click()
            expect(page.locator('#import-dialog')).not_to_be_visible()
            expect(page.locator('#nav-count')).to_have_text('100')
            expect(page.locator('#total-cartons')).to_have_text('182,500')
            # Reconnecting with the same public config must preserve the existing workspace.
            page.locator('#connection-settings').click();page.locator('#connect-submit').click()
            expect(page.locator('#setup-screen')).not_to_be_visible()
            expect(page.locator('#nav-count')).to_have_text('100')
            assert page.url==shared_link
            page.locator('#toast').evaluate('(el)=>el.hidden=true')
            page.screenshot(path=OUTPUT/'free-dashboard.png',full_page=True)
            other=browser.new_context(viewport={'width':1280,'height':900});wire(other)
            second=other.new_page();second.goto(shared_link,wait_until='networkidle')
            expect(second.locator('#nav-count')).to_have_text('100')
            expect(second.locator('#local-demo-note')).not_to_be_visible()

            # Updates require consent and can be followed from another browser.
            updated=sample_rows(100);updated[0]['delivered']='2026-10-17'
            upload(second,updated,'updated.xlsx')
            expect(second.locator('#accept-updates')).to_be_visible(timeout=15000)
            expect(second.locator('#confirm-import')).to_be_disabled()
            second.locator('#accept-updates').check();second.locator('#confirm-import').click()
            expect(second.locator('#import-dialog')).not_to_be_visible()
            page.locator('#refresh').click()
            page.locator('[data-view=records]').click();page.locator('#search').fill('DEMO-0001')
            expect(page.locator('#records-table tbody tr')).to_have_count(2)
            expect(page.locator('#records-table')).to_contain_text('17 ต.ค.')
            page.locator('#records-table tbody tr').first.click()
            expect(page.locator('#detail-body')).to_contain_text('วันที่รับเอกสารแก้ไขล่าสุด')
            page.locator('#detail-dialog .close-button').click()
            for kind in ['xlsx','csv']:
                page.locator('#export-toggle').click()
                with page.expect_download() as dl:page.locator(f'[data-export={kind}]').click()
                path=Path(directory)/('report.'+kind);dl.value.save_as(path)
                if kind=='xlsx':
                    wb=load_workbook(path);assert wb['Filtered records'].max_row==3;assert wb['Job totals'].max_row==2
                else:assert path.read_bytes().startswith(b'\xef\xbb\xbf')
            page.locator('#clear-filters').click();page.locator('#next-page').click();expect(page.locator('#page-summary')).to_contain_text('21–40')
            page.locator('#more-filters').click();page.locator('input[data-field=received_latest]').fill('2026-09-28')
            expect(page.locator('#result-label')).to_have_text('0 รายการ')
            page.locator('#clear-filters').click();page.locator('#more-filters').click()
            page.locator('[data-view=history]').click();expect(page.locator('#history-table tbody tr')).to_have_count(2)
            page.locator('#history-table tbody tr').first.click();expect(page.locator('#detail-body')).to_contain_text('2026-10-17')
            page.locator('#detail-dialog .close-button').click();page.locator('[data-view=dashboard]').click()
            page.evaluate('window.print = () => {window.printRequested = true;}')
            page.locator('#export-toggle').click();page.locator('[data-export=pdf]').click();assert page.evaluate('window.printRequested===true')
            page.emulate_media(media='print');pdf=page.pdf(path=OUTPUT/'free-report.pdf',format='A4',print_background=True);assert pdf.startswith(b'%PDF');page.emulate_media(media='screen')
            for width in [390,768,1440]:
                page.set_viewport_size({'width':width,'height':900})
                assert page.evaluate('document.documentElement.scrollWidth<=document.documentElement.clientWidth'),f'Layout overflow at {width}'
            page.set_viewport_size({'width':390,'height':844});page.locator('#toast').evaluate('(el)=>el.hidden=true');page.screenshot(path=OUTPUT/'free-mobile.png',full_page=True)
            page.reload(wait_until='networkidle');expect(page.locator('#nav-count')).to_have_text('100')

            # A different workspace link must not expose this team's records.
            isolated=browser.new_context();wire(isolated);third=isolated.new_page();third.goto(base,wait_until='networkidle')
            third.locator('#supabase-url').fill('https://harbor-test.supabase.co');third.locator('#supabase-key').fill('sb_publishable_test_key_for_browser_only');third.locator('#connect-submit').click();expect(third.locator('#nav-count')).to_have_text('0')
            assert errors==[],errors
            browser.close()
    print('PASS: static project-path site, real PostgreSQL RPCs as anon, 2-browser shared data/update, workspace isolation, consent, filters, details, pagination, history, Excel/CSV/PDF, reload persistence, demo separation and responsive layouts. Hosted Supabase and public Pages activation still require user setup.')


try:
    if __name__=='__main__':run()
finally:
    if created_workspaces:
        with psycopg.connect(DSN) as conn:
            for workspace in created_workspaces:
                for table in ['revisions','records','batches','workspaces']:
                    conn.execute('delete from harbor_private.'+table+' where '+('id' if table=='workspaces' else 'workspace_id')+'=%s',[workspace])
