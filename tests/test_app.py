import csv
import io
import os
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

TEST_DIR = tempfile.TemporaryDirectory(prefix='harbor-tests-')
os.environ['DATA_DIR'] = TEST_DIR.name

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app, db
from tests.fixtures import sample_rows, workbook


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        with db() as conn:
            for table in ['records','batches','revisions','stages']:
                conn.execute('DELETE FROM '+table)
            conn.commit()

    def preview(self, rows=None, content=None, extra=None):
        return self.client.post('/api/import/preview', files={'file':('test.xlsx',content or workbook(rows,extra),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})

    def commit(self, preview, **kwargs):
        return self.client.post('/api/import/confirm',json={'token':preview.json()['token'],**kwargs})

    def seed(self, rows=None):
        p=self.preview(rows)
        self.assertEqual(p.status_code,200,p.text)
        r=self.commit(p)
        self.assertEqual(r.status_code,200,r.text)
        return p

    def test_counts_job_totals_and_split_dates(self):
        self.seed()
        c=self.client.get('/api/dashboard').json()['counts']
        self.assertEqual((c['records'],c['jobs'],c['invoices'],c['containers']),(12,6,6,12))
        self.assertEqual(c['cartons'],sum(1200+i*100 for i in range(6)))
        self.assertEqual(c['cbm'],sum(64.5+i*5 for i in range(6)))
        self.assertEqual((c['pending'],c['delivered'],c['late'],c['on_time']),(4,8,4,4))
        r=self.client.get('/api/records').json()['rows'][0]
        self.assertEqual(r['received_first'],'2026-09-25T09:00:00')
        self.assertEqual(r['received_latest'],'2026-09-27T14:30:00')
        self.assertNotIn('declaration_status',r)

    def test_filters_and_pagination(self):
        self.seed(sample_rows(40))
        first=self.client.get('/api/records?page=1').json()
        second=self.client.get('/api/records?page=2').json()
        self.assertEqual((first['total'],len(first['rows']),len(second['rows'])),(40,20,20))
        self.assertFalse({r['container'] for r in first['rows']} & {r['container'] for r in second['rows']})
        self.assertEqual(self.client.get('/api/records?q=demo-0001').json()['total'],2)
        self.assertEqual(self.client.get('/api/records?date_field=received_latest&from=2026-09-27&to=2026-09-27').json()['total'],40)
        self.assertEqual(self.client.get('/api/records?date_field=received_first&from=2026-09-27').json()['total'],0)
        self.assertEqual(self.client.get('/api/records?f_cartons=1200').json()['total'],2)
        self.assertEqual(self.client.get('/api/records?from=2026-10-09&to=2026-10-01').status_code,400)

    def test_same_file_and_unchanged_records(self):
        content=workbook()
        p=self.preview(content=content);self.assertEqual(self.commit(p).status_code,200)
        self.assertEqual(self.preview(content=content).status_code,409)
        rows=sample_rows();rows.reverse()
        p=self.preview(rows);self.assertEqual(p.json()['counts']['unchanged'],12)
        self.assertEqual(self.commit(p).status_code,200)
        self.assertEqual(self.client.get('/api/meta').json()['total'],12)

    def test_updates_need_confirmation_and_keep_history(self):
        self.seed()
        rows=sample_rows();rows[0]['delivered']=datetime(2026,10,9)
        p=self.preview(rows);self.assertEqual(p.json()['counts']['updated'],1)
        self.assertEqual(self.commit(p).status_code,400)
        result=self.commit(p,accept_updates=True)
        self.assertEqual(result.status_code,200,result.text)
        revisions=self.client.get('/api/history/'+result.json()['batch_id']).json()['rows']
        self.assertEqual(len(revisions),1)
        self.assertEqual(revisions[0]['before']['delivered'],'')
        self.assertTrue(revisions[0]['after']['delivered'].startswith('2026-10-09'))

    def test_stale_preview_cannot_overwrite_concurrent_change(self):
        self.seed()
        a=sample_rows();a[0]['delivered']=datetime(2026,10,9)
        b=sample_rows();b[0]['delivered']=datetime(2026,10,10)
        pa=self.preview(a);pb=self.preview(b)
        self.assertEqual(self.commit(pa,accept_updates=True).status_code,200)
        self.assertEqual(self.commit(pb,accept_updates=True).status_code,409)

    def test_partial_invalid_rows_require_opt_in(self):
        rows=sample_rows();rows.append({'port':'stray cell'})
        p=self.preview(rows);self.assertEqual(p.json()['counts']['invalid'],1)
        self.assertEqual(self.commit(p).status_code,400)
        self.assertEqual(self.commit(p,skip_invalid=True).status_code,200)
        self.assertEqual(self.client.get('/api/meta').json()['total'],12)

    def test_inconsistent_job_totals_atomic_rejection(self):
        rows=sample_rows();rows[1]['cartons']=999
        p=self.preview(rows);self.assertEqual(self.commit(p).status_code,422)
        self.assertEqual(self.client.get('/api/meta').json()['total'],0)

    def test_conflicting_keys_inside_file_reject_whole_preview(self):
        rows=sample_rows();rows.append(rows[0]|{'carrier':'Other'})
        self.assertEqual(self.preview(rows).status_code,422)

    def test_invalid_and_formula_files(self):
        self.assertEqual(self.preview(content=b'not a workbook').status_code,422)
        rows=sample_rows();rows[0]['job']='=1+1'
        p=self.preview(rows);self.assertEqual(p.status_code,200);self.assertEqual(p.json()['counts']['invalid'],1)
        rows=sample_rows();rows[0]['cartons']=-1
        self.assertEqual(self.preview(rows).json()['counts']['invalid'],1)

    def test_exports_filter_and_protect_cells(self):
        rows=sample_rows();rows[0]['seal']='@SUM(1,1)'
        self.seed(rows)
        response=self.client.get('/api/export/xlsx?q=DEMO-0001')
        self.assertEqual(response.status_code,200)
        wb=load_workbook(io.BytesIO(response.content),data_only=False)
        self.assertEqual(wb['Filtered records'].max_row,3)
        self.assertEqual(wb['Job totals'].max_row,2)
        self.assertEqual(wb['Job totals'].cell(2,2).value,1200)
        cells=list(wb['Filtered records'].values)
        seal=cells[1][cells[0].index('Seal')]
        self.assertEqual(seal,"'@SUM(1,1)")
        response=self.client.get('/api/export/csv?q=DEMO-0001')
        out=list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(len(out),3)
        self.assertIn('วันที่รับครั้งแรก',out[0])

    def test_extra_header_filter_and_missing_header_error(self):
        rows=sample_rows();rows[0]['account']='Demo account'
        p=self.preview(rows,extra=[('account','Customer Account')]);self.assertEqual(self.commit(p).status_code,200)
        result=self.client.get('/api/records',params={'f_extra:Customer Account':'Demo'})
        self.assertEqual(result.status_code,200);self.assertEqual(result.json()['total'],1)
        wb=load_workbook(io.BytesIO(workbook()));wb.active.cell(1,2).value='Not ETA';out=io.BytesIO();wb.save(out)
        self.assertEqual(self.preview(content=out.getvalue()).status_code,422)

    def test_30_concurrent_readers(self):
        self.seed()
        def read(_):
            with TestClient(app) as c:
                return c.get('/api/dashboard').json()['counts']['records']
        with ThreadPoolExecutor(max_workers=30) as pool:
            self.assertEqual(list(pool.map(read,range(30))),[12]*30)

    def test_health_and_cross_site_write(self):
        self.assertEqual(self.client.get('/api/health').json()['status'],'ok')
        r=self.client.post('/api/import/confirm',json={'token':'fake'},headers={'Origin':'https://unrelated.example'})
        self.assertEqual(r.status_code,403)


if __name__=='__main__':
    unittest.main()
