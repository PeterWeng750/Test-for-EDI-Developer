"""End-to-end tests with synthetic data, a temporary database and a real browser.

CHROMIUM_PATH=/usr/bin/chromium python tests/browser_smoke.py
"""
import io
import os
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from urllib.request import urlopen

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook
from tests.fixtures import workbook, sample_rows

OUTPUT = Path(os.environ.get('OUTPUT_DIR', ROOT / 'test-results'))
OUTPUT.mkdir(parents=True, exist_ok=True)


def run():
    with tempfile.TemporaryDirectory(prefix='harbor-browser-') as directory:
        with socket.socket() as sock:
            sock.bind(('127.0.0.1', 0))
            port = sock.getsockname()[1]
        base = f'http://127.0.0.1:{port}'
        env = os.environ | {'DATA_DIR': directory}
        command = [os.environ.get('SERVER_PYTHON', sys.executable), '-m', 'uvicorn', 'app.main:app', '--host', '127.0.0.1', '--port', str(port)]

        def start():
            process = subprocess.Popen(command, cwd=ROOT, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            deadline = time.monotonic()+15
            while time.monotonic()<deadline:
                try:
                    with urlopen(base+'/api/health',timeout=1) as response:
                        assert response.status == 200
                    return process
                except OSError:
                    if process.poll() is not None:
                        raise RuntimeError('Server exited before startup')
                    time.sleep(.1)
            process.terminate()
            raise RuntimeError('Server startup timed out')

        server = start()
        try:
            with sync_playwright() as p:
                launch = {'headless': True, 'args':['--no-sandbox']}
                if os.environ.get('CHROMIUM_PATH'):
                    launch['executable_path'] = os.environ['CHROMIUM_PATH']
                browser = p.chromium.launch(**launch)
                page = browser.new_page(viewport={'width':1440,'height':1000})
                errors=[]
                page.on('pageerror',lambda error: errors.append(str(error)))
                page.goto(base,wait_until='networkidle')
                expect(page.locator('#empty-state')).to_be_visible()
                page.screenshot(path=OUTPUT/'empty-state.png',full_page=True)
                rows=sample_rows(100)
                rows.append({'port':'Invalid synthetic row'})
                page.locator('[data-import]').first.click()
                page.locator('#file-input').set_input_files({'name':'synthetic-demo.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':workbook(rows)})
                expect(page.locator('#confirm-import')).to_be_visible()
                expect(page.locator('#confirm-import')).to_be_disabled()
                page.locator('#accept-invalid').check()
                page.locator('#confirm-import').click()
                expect(page.locator('#import-dialog')).not_to_be_visible()
                expect(page.locator('#nav-count')).to_have_text('100')
                expect(page.locator('#kpi-grid')).to_contain_text('50')
                page.locator('#toast').evaluate('(el)=>el.hidden=true')
                page.screenshot(path=OUTPUT/'dashboard.png',full_page=True)

                # Another browser context reads the same persisted server-side data.
                second=browser.new_page()
                second.goto(base,wait_until='networkidle')
                expect(second.locator('#nav-count')).to_have_text('100')
                second.close()

                page.locator('[data-view=records]').click()
                expect(page.locator('#records-table tbody tr')).to_have_count(20)
                page.locator('#search').fill('DEMO-0001')
                expect(page.locator('#records-table tbody tr')).to_have_count(2)
                page.locator('#records-table tbody tr').first.click()
                expect(page.locator('#detail-body')).to_contain_text('วันที่รับครั้งแรก')
                page.locator('#detail-dialog .close-button').click()
                page.locator('#export-toggle').click()
                with page.expect_download() as dl:
                    page.locator('[data-export=xlsx]').click()
                path=Path(directory)/'filtered.xlsx'
                dl.value.save_as(path)
                report=load_workbook(path)
                assert report['Filtered records'].max_row == 3
                assert report['Job totals'].max_row == 2
                page.locator('#export-toggle').click()
                with page.expect_download() as dl:
                    page.locator('[data-export=csv]').click()
                csv_path=Path(directory)/'filtered.csv'
                dl.value.save_as(csv_path)
                assert csv_path.read_bytes().startswith(b'\xef\xbb\xbf')
                page.locator('#clear-filters').click()
                expect(page.locator('#records-table tbody tr')).to_have_count(20)
                page.locator('#next-page').click()
                expect(page.locator('#page-summary')).to_contain_text('21–40')

                page.locator('#more-filters').click()
                page.locator('input[data-field=received_first]').fill('2026-09-25')
                expect(page.locator('#result-label')).to_have_text('100 รายการ')
                page.locator('input[data-field=received_latest]').fill('2026-09-28')
                expect(page.locator('#result-label')).to_have_text('0 รายการ')
                page.locator('#clear-filters').click()
                page.locator('#more-filters').click()

                page.locator('[data-view=history]').click()
                expect(page.locator('#history-table tbody tr')).to_have_count(1)
                page.locator('#history-table tbody tr').click()
                expect(page.locator('#detail-dialog')).to_be_visible()
                page.locator('#detail-dialog .close-button').click()
                page.locator('[data-view=dashboard]').click()
                page.evaluate('window.print = () => { window.printRequested = true; }')
                page.locator('#export-toggle').click()
                page.locator('[data-export=pdf]').click()
                assert page.evaluate('window.printRequested === true')
                page.emulate_media(media='print')
                pdf=page.pdf(path=OUTPUT/'sample-report.pdf',format='A4',print_background=True)
                assert pdf.startswith(b'%PDF') and len(pdf)>10000
                page.emulate_media(media='screen')
                for width in [390,768,1440]:
                    page.set_viewport_size({'width':width,'height':900})
                    assert page.evaluate('document.documentElement.scrollWidth <= document.documentElement.clientWidth'), f'Layout overflow at {width}px'
                page.set_viewport_size({'width':390,'height':844})
                page.locator('#toast').evaluate('(el)=>el.hidden=true')
                page.screenshot(path=OUTPUT/'mobile.png',full_page=True)

                server.terminate();server.wait(timeout=10)
                server=start()
                page.reload(wait_until='networkidle')
                expect(page.locator('#nav-count')).to_have_text('100')
                assert errors==[], errors
                browser.close()
                print('PASS: browser import/invalid-row consent, shared data, filters, details, pagination, history, Excel/CSV/PDF, 390/768/1440px layout, restart persistence; no JavaScript errors.')
        finally:
            server.terminate()
            server.wait(timeout=10)


if __name__ == '__main__':
    run()
