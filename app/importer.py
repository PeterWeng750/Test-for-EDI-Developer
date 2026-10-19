"""Workbook parsing. Uploaded cell contents are data, never instructions."""
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

FIELDS = [
    ('job', 'Job I'), ('eta', 'ETA'), ('ata', 'ATA'), ('invoice', 'Invoice No'),
    ('origin', 'Origin(Port Name)'), ('port', 'Release Port'), ('carrier', 'Carrier Name'),
    ('container', 'CTNR NO.'), ('container_count', 'CTNR NO. Count'), ('seal', 'Seal'),
    ('size', 'Container Size'), ('cartons', 'Cartons'), ('cbm', 'CBM'), ('file_date', 'FileNameDate'),
    ('do_release', 'DO Release'), ('received', 'Shipping Document receive (วันที่ได้เอกสารในการทำใบขน)'),
    ('issue', 'Shipping Document issue reason (ข้อมูลผิดรอแก้ไข)'),
    ('complete', 'Shipping Document Complete (วันที่ได้เอกสารพร้อมทำใบขน)'),
    ('declaration_status', 'Status Declaration'),
    ('customs', 'Customs Relese Import (Declaration complete) (0209)'),
    ('plan', 'Delivery Plan'), ('release', 'Import Decleration Relese (0309)'),
    ('delivered', 'Delivery Date (Customs)'), ('place', 'Delivery Place'),
]
LABELS = dict(FIELDS) | {'received_first': 'วันที่รับครั้งแรก', 'received_latest': 'วันที่รับเอกสารแก้ไขล่าสุด'}
DATE_FIELDS = {'eta', 'ata', 'file_date', 'do_release', 'complete', 'customs', 'plan', 'release', 'delivered', 'received_first', 'received_latest'}
NUM_FIELDS = {'container_count', 'cartons', 'cbm'}
KEY_FIELDS = ('job', 'invoice', 'container')
MAX_ROWS = 10000


def norm(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().casefold()


def digest(record):
    return hashlib.sha256(json.dumps(record, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def key(record):
    return json.dumps([norm(record[k]) for k in KEY_FIELDS], ensure_ascii=False)


def parse_date(value, epoch):
    if value is None or value == '':
        return ''
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(ZoneInfo('Asia/Bangkok')).replace(tzinfo=None)
        return value.isoformat(timespec='seconds')
    if isinstance(value, date):
        return value.isoformat() + 'T00:00:00'
    text = str(value).strip()
    if re.fullmatch(r'\d{8}', text):
        return datetime.strptime(text, '%Y%m%d').isoformat()
    if isinstance(value, (int, float)):
        parsed = from_excel(value, epoch)
        if not isinstance(parsed, datetime):
            raise ValueError('คอลัมน์วันที่ต้องมีวันที่ ไม่ใช่เฉพาะเวลา')
        return parsed.isoformat()
    try:
        return parse_date(datetime.fromisoformat(text), epoch)
    except ValueError:
        for fmt in ('%d/%m/%Y', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M'):
            try:
                return datetime.strptime(text, fmt).isoformat(timespec='seconds')
            except ValueError:
                pass
    raise ValueError('รูปแบบวันที่ไม่ถูกต้อง (ใช้ YYYY-MM-DD หรือ DD/MM/YYYY)')


def parse_workbook(content):
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ValueError('กรุณาเลือกไฟล์ Excel .xlsx ที่ไม่เข้ารหัส')
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
            raise ValueError('ไฟล์ Excel หลังคลายขนาดใหญ่เกินขีดจำกัดรุ่นทดลอง 80 MB')
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    aliases = {norm(label): name for name, label in LABELS.items()}
    aliases.update({'po': 'po', 'po no': 'po', 'po no.': 'po', 'po number': 'po'})
    output, errors, sheets, duplicates = [], [], [], 0
    seen, total = {}, 0
    try:
        for ws in wb:
            iterator = ws.iter_rows()
            first = next(iterator, ())
            mapping = {i: aliases.get(norm(c.value), 'extra:' + str(c.value).strip()) for i, c in enumerate(first) if c.value is not None}
            if not set(KEY_FIELDS).issubset(mapping.values()):
                continue
            required = {k for k, _ in FIELDS} - {'declaration_status'}
            missing_headers = required - set(mapping.values())
            if missing_headers:
                raise ValueError(f'ชีต {ws.title}: ขาด Header ' + ', '.join(LABELS[k] for k in sorted(missing_headers)))
            if len(mapping.values()) != len(set(mapping.values())):
                raise ValueError(f'ชีต {ws.title}: Header ซ้ำ กรุณาตรวจสอบ')
            sheets.append(ws.title)
            for line, cells in enumerate(iterator, 2):
                if not any(c.value is not None for c in cells):
                    continue
                total += 1
                if total > MAX_ROWS:
                    raise ValueError('รุ่นทดลองรองรับไม่เกิน 10,000 แถวต่อไฟล์')
                raw = {name: cells[i].value if i < len(cells) else None for i, name in mapping.items()}
                try:
                    if any(i < len(cells) and cells[i].data_type == 'f' for i in mapping):
                        raise ValueError('พบสูตร Excel กรุณาวางเป็นค่าก่อนนำเข้า')
                    missing = [LABELS[k] for k in KEY_FIELDS if raw.get(k) is None or not str(raw[k]).strip()]
                    if missing:
                        raise ValueError('ขาด ' + ', '.join(missing))
                    record = {}
                    for name, value in raw.items():
                        if name == 'declaration_status':
                            continue
                        if name in DATE_FIELDS:
                            record[name] = parse_date(value, wb.epoch)
                        elif name in NUM_FIELDS:
                            n = None if value in (None, '') else float(str(value).replace(',', ''))
                            if n is not None and (not math.isfinite(n) or n < 0 or (name != 'cbm' and n != int(n))):
                                raise ValueError(f'{LABELS[name]} ต้องเป็นจำนวนที่ไม่ติดลบ' + ('และเป็นจำนวนเต็ม' if name != 'cbm' else ''))
                            record[name] = n
                        else:
                            record[name] = '' if value is None else str(value).strip()
                    times = re.findall(r'\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}', record.get('received', ''))
                    if record.get('received') and len(times) not in (1, 2):
                        raise ValueError('ช่องรับเอกสารต้องมีวันเวลา 1 หรือ 2 ชุด รูปแบบ YYYY-MM-DD HH:MM:SS')
                    record['received_first'] = parse_date(times[0], wb.epoch) if times else ''
                    record['received_latest'] = parse_date(times[-1], wb.epoch) if len(times) == 2 else ''
                    if record['received_latest'] and record['received_latest'] < record['received_first']:
                        raise ValueError('วันรับเอกสารแก้ไขล่าสุดอยู่ก่อนวันรับครั้งแรก')
                    identity = key(record)
                    if identity in seen:
                        if seen[identity] != record:
                            raise RuntimeError('คีย์ Job + Invoice + ตู้ ซ้ำแต่ข้อมูลต่างกันภายในไฟล์ กรุณาแก้ไขไฟล์ก่อนนำเข้า')
                        duplicates += 1
                        continue
                    seen[identity] = record
                    output.append({'key': identity, 'record': record, 'row': line, 'sheet': ws.title})
                except (ValueError, TypeError, OverflowError) as exc:
                    errors.append({'sheet': ws.title, 'row': line, 'message': str(exc)})
        if not sheets:
            raise ValueError('ไม่พบชีตที่มี Header: Job I, Invoice No และ CTNR NO.')
        return {'rows': output, 'errors': errors, 'sheets': sheets, 'duplicates_in_file': duplicates, 'total': total}
    finally:
        wb.close()
