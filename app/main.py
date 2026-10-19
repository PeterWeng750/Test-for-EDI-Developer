import csv
import hashlib
import io
import json
import os
import secrets
import sqlite3
import threading
import zipfile
from xml.etree.ElementTree import ParseError
from collections import defaultdict
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from starlette.concurrency import run_in_threadpool

from .importer import DATE_FIELDS, FIELDS, LABELS, NUM_FIELDS, digest, norm, parse_workbook

DATA_DIR = Path(os.environ.get('DATA_DIR', './data')).resolve()
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = DATA_DIR / 'dashboard.sqlite3'
TZ = ZoneInfo('Asia/Bangkok')
MAX_FILE = 10 * 1024 * 1024
IMPORT_SLOTS = threading.BoundedSemaphore(2)
app = FastAPI(title='Harbor — Import Operations', docs_url=None, redoc_url=None)


@contextmanager
def db():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()


with db() as conn:
    conn.executescript('''
    PRAGMA journal_mode=WAL;
    CREATE TABLE IF NOT EXISTS records (identity TEXT PRIMARY KEY, data TEXT NOT NULL, hash TEXT NOT NULL, updated_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS batches (id TEXT PRIMARY KEY, filename TEXT NOT NULL, file_hash TEXT UNIQUE NOT NULL, created_at TEXT NOT NULL, added INTEGER, updated INTEGER, skipped INTEGER, invalid INTEGER);
    CREATE TABLE IF NOT EXISTS stages (id TEXT PRIMARY KEY, content TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS revisions (id INTEGER PRIMARY KEY, identity TEXT NOT NULL, batch_id TEXT NOT NULL, old_data TEXT, new_data TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE INDEX IF NOT EXISTS idx_eta ON records(json_extract(data, '$.eta'));
    CREATE INDEX IF NOT EXISTS idx_plan ON records(json_extract(data, '$.plan'));
    CREATE INDEX IF NOT EXISTS idx_job ON records(json_extract(data, '$.job'));
    ''')


@app.middleware('http')
async def headers(request, call_next):
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        if request.headers.get('sec-fetch-site') == 'cross-site':
            return Response('Cross-site writes are not allowed', status_code=403)
        origin = request.headers.get('origin')
        if origin:
            from urllib.parse import urlparse
            if urlparse(origin).netloc != request.headers.get('host'):
                return Response('Origin mismatch', status_code=403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'"
    if request.url.path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-store'
    return response


def now():
    return datetime.now(TZ).isoformat(timespec='seconds')


def label(field):
    return LABELS.get(field, 'PO No' if field == 'po' else field.removeprefix('extra:'))


def all_fields(conn):
    fields = [k for k, _ in FIELDS if k not in ('received', 'declaration_status')]
    fields[fields.index('complete'):fields.index('complete')] = ['received_first', 'received_latest']
    extras = conn.execute("SELECT DISTINCT j.key FROM records, json_each(records.data) j WHERE j.key = 'po' OR j.key LIKE 'extra:%'").fetchall()
    return fields + sorted(r[0] for r in extras)


def query(conn, params):
    clauses, values = [], []
    fields = all_fields(conn)
    search = params.get('q', '').strip().lower()
    if search:
        clauses.append('(' + ' OR '.join('instr(lower(CAST(json_extract(data, ?) AS TEXT)), ?) > 0' for _ in ['job', 'invoice', 'container', 'po', 'carrier', 'origin', 'place']) + ')')
        for k in ['job', 'invoice', 'container', 'po', 'carrier', 'origin', 'place']:
            values.extend(['$.' + k, search])
    date_field = params.get('date_field', 'eta')
    if date_field not in DATE_FIELDS:
        raise HTTPException(400, 'คอลัมน์วันที่ไม่ถูกต้อง')
    start, end = params.get('from', ''), params.get('to', '')
    for val in [start, end]:
        if val:
            try:
                datetime.strptime(val, '%Y-%m-%d')
            except ValueError:
                raise HTTPException(400, 'ช่วงวันที่ไม่ถูกต้อง')
    if start and end and start > end:
        raise HTTPException(400, 'วันเริ่มต้นต้องไม่อยู่หลังวันสิ้นสุด')
    for value, operator in [(start, '>='), (end, '<=')]:
        if value:
            clauses.append(f"substr(json_extract(data, ?), 1, 10) {operator} ? AND json_extract(data, ?) != ''")
            values.extend(['$.' + date_field, value, '$.' + date_field])
    for k, value in params.items():
        if k.startswith('f_') and value:
            name = k[2:]
            if name not in fields:
                raise HTTPException(400, 'ไม่พบคอลัมน์ตัวกรอง')
            path = '$.' + json.dumps(name, ensure_ascii=False)
            if name in NUM_FIELDS:
                try:
                    num = float(value)
                except ValueError:
                    raise HTTPException(400, 'ตัวกรองจำนวนต้องเป็นตัวเลข')
                clauses.append('json_extract(data, ?) = ?')
                values.extend([path, num])
            else:
                clauses.append('instr(lower(CAST(json_extract(data, ?) AS TEXT)), ?) > 0')
                values.extend([path, value.lower()])
    status = params.get('status', '')
    delivered = "coalesce(json_extract(data, '$.delivered'), '')"
    plan = "coalesce(json_extract(data, '$.plan'), '')"
    if status == 'delivered':
        clauses.append(f"{delivered} != ''")
    elif status == 'pending':
        clauses.append(f"{delivered} = ''")
    elif status == 'overdue':
        clauses.append(f"{delivered} = '' AND {plan} != '' AND substr({plan},1,10) < ?")
        values.append(datetime.now(TZ).date().isoformat())
    elif status == 'late':
        clauses.append(f"{delivered} != '' AND {plan} != '' AND substr({delivered},1,10) > substr({plan},1,10)")
    elif status:
        raise HTTPException(400, 'สถานะไม่ถูกต้อง')
    return (' WHERE ' + ' AND '.join(clauses) if clauses else ''), values


def selected(conn, params):
    where, values = query(conn, params)
    return [json.loads(r['data']) for r in conn.execute('SELECT data FROM records' + where + ' ORDER BY updated_at DESC, identity', values)]


@app.get('/api/health')
def health():
    with db() as conn:
        conn.execute('SELECT 1').fetchone()
    return {'status': 'ok', 'timezone': 'Asia/Bangkok'}


@app.get('/api/meta')
def meta():
    with db() as conn:
        fields = all_fields(conn)
        options = {}
        for k in ['carrier', 'origin', 'port', 'size', 'place']:
            options[k] = [r[0] for r in conn.execute('SELECT DISTINCT json_extract(data, ?) AS value FROM records WHERE value IS NOT NULL AND value != ? ORDER BY value', ['$.' + k, ''])]
        total = conn.execute('SELECT count(*) FROM records').fetchone()[0]
    return {'fields': [{'key': k, 'label': label(k), 'type': 'date' if k in DATE_FIELDS else 'number' if k in NUM_FIELDS else 'text'} for k in fields], 'options': options, 'total': total}


def preview(content, filename):
    if not IMPORT_SLOTS.acquire(blocking=False):
        raise HTTPException(429, 'มีไฟล์กำลังประมวลผล กรุณาลองอีกครั้งในสักครู่')
    try:
        file_hash = hashlib.sha256(content).hexdigest()
        with db() as conn:
            if conn.execute('SELECT 1 FROM batches WHERE file_hash = ?', [file_hash]).fetchone():
                raise HTTPException(409, 'ไฟล์นี้ถูกนำเข้าแล้ว กรุณาเลือกไฟล์ใหม่หรือไฟล์ที่แก้ไขแล้ว')
        try:
            parsed = parse_workbook(content)
        except (ValueError, RuntimeError, KeyError, zipfile.BadZipFile, ParseError, OSError) as exc:
            raise HTTPException(422, str(exc))
        counts = {'new': 0, 'updated': 0, 'unchanged': parsed['duplicates_in_file'], 'invalid': len(parsed['errors'])}
        changes = []
        with db() as conn:
            for row in parsed['rows']:
                old = conn.execute('SELECT hash, data FROM records WHERE identity = ?', [row['key']]).fetchone()
                row['old_hash'] = old['hash'] if old else None
                row['action'] = 'new' if not old else 'unchanged' if old['hash'] == digest(row['record']) else 'updated'
                counts[row['action']] += 1
                if row['action'] == 'updated' and len(changes) < 100:
                    before = json.loads(old['data'])
                    changed_fields = [{'field': label(k), 'before': before.get(k), 'after': row['record'].get(k)} for k in sorted(set(before) | set(row['record'])) if before.get(k) != row['record'].get(k)]
                    changes.append({'job': row['record']['job'], 'invoice': row['record']['invoice'], 'container': row['record']['container'], 'fields': changed_fields})
            token = secrets.token_urlsafe(24)
            stage = parsed | {'filename': filename, 'file_hash': file_hash, 'counts': counts}
            conn.execute('DELETE FROM stages WHERE created_at < ?', [(datetime.now(TZ) - timedelta(hours=24)).isoformat()])
            if conn.execute('SELECT count(*) FROM stages').fetchone()[0] >= 50:
                raise HTTPException(429, 'มีรายการรอยืนยันจำนวนมาก กรุณายืนยันรายการเดิมหรือลองใหม่ภายหลัง')
            conn.execute('INSERT INTO stages VALUES (?, ?, ?)', [token, json.dumps(stage, ensure_ascii=False), now()])
            conn.commit()
        return {'token': token, 'filename': filename, 'counts': counts, 'errors': parsed['errors'][:100], 'sheets': parsed['sheets'], 'changes': changes, 'sample': [r['record'] for r in parsed['rows'][:5]], 'total': parsed['total']}
    finally:
        IMPORT_SLOTS.release()


@app.post('/api/import/preview')
async def upload(file: UploadFile = File(...)):
    try:
        if not (file.filename or '').lower().endswith('.xlsx'):
            raise HTTPException(422, 'รุ่นทดลองรองรับไฟล์ .xlsx กรุณาบันทึก .xls เป็น .xlsx ก่อน')
        content = await file.read(MAX_FILE + 1)
        if len(content) > MAX_FILE:
            raise HTTPException(413, 'รุ่นทดลองรองรับไฟล์ไม่เกิน 10 MB')
        return await run_in_threadpool(preview, content, Path(file.filename).name)
    finally:
        await file.close()


class Confirm(BaseModel):
    token: str
    accept_updates: bool = False
    skip_invalid: bool = False


@app.post('/api/import/confirm')
def confirm(body: Confirm):
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        row = conn.execute('SELECT content, created_at FROM stages WHERE id = ?', [body.token]).fetchone()
        if not row or datetime.fromisoformat(row['created_at']) < datetime.now(TZ) - timedelta(hours=24):
            raise HTTPException(410, 'รายการตรวจสอบหมดอายุ กรุณาเลือกไฟล์อีกครั้ง')
        stage = json.loads(row['content'])
        counts = stage['counts']
        if counts['updated'] and not body.accept_updates:
            raise HTTPException(400, 'กรุณายืนยันการอัปเดตรายการเดิม')
        if counts['invalid'] and not body.skip_invalid:
            raise HTTPException(400, 'กรุณายืนยันข้ามแถวที่ไม่สมบูรณ์')
        if not stage['rows']:
            raise HTTPException(422, 'ไม่มีข้อมูลที่นำเข้าได้')
        if conn.execute('SELECT 1 FROM batches WHERE file_hash = ?', [stage['file_hash']]).fetchone():
            raise HTTPException(409, 'ไฟล์นี้ถูกนำเข้าแล้ว')
        if conn.execute('SELECT count(*) FROM records').fetchone()[0] + counts['new'] > 100000:
            raise HTTPException(422, 'รุ่นทดลองรองรับข้อมูลสะสม 100,000 รายการ กรุณาขยายระบบก่อนเพิ่มข้อมูล')
        for item in stage['rows']:
            current = conn.execute('SELECT hash FROM records WHERE identity = ?', [item['key']]).fetchone()
            if (current['hash'] if current else None) != item['old_hash']:
                raise HTTPException(409, 'ข้อมูลเปลี่ยนโดยผู้ใช้อื่น กรุณาเลือกไฟล์ใหม่เพื่อตรวจสอบอีกครั้ง')
        prospective = {r['identity']: json.loads(r['data']) for r in conn.execute('SELECT identity, data FROM records')}
        prospective.update({r['key']: r['record'] for r in stage['rows']})
        jobs = {}
        for item in prospective.values():
            job = norm(item['job'])
            totals = (item.get('cartons'), item.get('cbm'))
            if job in jobs and jobs[job] != totals:
                raise HTTPException(422, 'Cartons/CBM ของ Job เดียวกันไม่ตรงกัน กรุณาแก้ไขและอัปโหลดข้อมูลครบทุกแถวของ Job นั้น')
            jobs[job] = totals
        timestamp, batch_id = now(), secrets.token_urlsafe(12)
        for item in stage['rows']:
            if item['action'] == 'unchanged':
                continue
            previous = conn.execute('SELECT data FROM records WHERE identity = ?', [item['key']]).fetchone()
            serialized = json.dumps(item['record'], ensure_ascii=False)
            conn.execute('INSERT INTO revisions(identity,batch_id,old_data,new_data,created_at) VALUES(?,?,?,?,?)', [item['key'], batch_id, previous['data'] if previous else None, serialized, timestamp])
            conn.execute('INSERT INTO records VALUES(?,?,?,?) ON CONFLICT(identity) DO UPDATE SET data=excluded.data, hash=excluded.hash, updated_at=excluded.updated_at', [item['key'], serialized, digest(item['record']), timestamp])
        conn.execute('INSERT INTO batches VALUES(?,?,?,?,?,?,?,?)', [batch_id, stage['filename'], stage['file_hash'], timestamp, counts['new'], counts['updated'], counts['unchanged'], counts['invalid']])
        conn.execute('DELETE FROM stages WHERE id = ?', [body.token])
        conn.commit()
    return {'batch_id': batch_id, 'counts': counts}


@app.get('/api/records')
def records(request: Request, page: int = 1, page_size: int = 20):
    if page < 1 or page_size not in (20, 50, 100):
        raise HTTPException(400, 'การแบ่งหน้าไม่ถูกต้อง')
    with db() as conn:
        where, values = query(conn, request.query_params)
        total = conn.execute('SELECT count(*) FROM records' + where, values).fetchone()[0]
        rows = conn.execute('SELECT data, updated_at FROM records' + where + ' ORDER BY updated_at DESC, identity LIMIT ? OFFSET ?', values + [page_size, (page - 1) * page_size]).fetchall()
    return {'total': total, 'page': page, 'page_size': page_size, 'rows': [json.loads(r['data']) | {'updated_at': r['updated_at']} for r in rows]}


@app.get('/api/dashboard')
def dashboard(request: Request):
    with db() as conn:
        rows = selected(conn, request.query_params)
        last = conn.execute('SELECT max(created_at) FROM batches').fetchone()[0]
    today = datetime.now(TZ).date().isoformat()
    jobs = {norm(r['job']): r for r in rows}
    counts = {'jobs': len(jobs), 'invoices': len({norm(r['invoice']) for r in rows}), 'containers': len({norm(r['container']) for r in rows}), 'records': len(rows), 'cartons': sum(r.get('cartons') or 0 for r in jobs.values()), 'cbm': round(sum(r.get('cbm') or 0 for r in jobs.values()), 3)}
    counts.update({'pending': 0, 'overdue': 0, 'delivered': 0, 'late': 0, 'on_time': 0, 'unknown_plan': 0})
    carriers, trends, issues = defaultdict(set), defaultdict(set), set()
    lead = defaultdict(list)
    attention = []
    for r in rows:
        plan, delivered = (r.get('plan') or '')[:10], (r.get('delivered') or '')[:10]
        if delivered:
            counts['delivered'] += 1
            counts['late' if delivered > plan else 'on_time'] += 1 if plan else 0
            counts['unknown_plan'] += int(not plan)
        else:
            counts['pending'] += 1
            if plan and plan < today:
                counts['overdue'] += 1
                attention.append(r)
        carriers[r.get('carrier') or 'ไม่ระบุ'].add(norm(r['container']))
        if r.get('ata'):
            trends[r['ata'][:10]].add(norm(r['job']))
        if r.get('issue'):
            issues.add(norm(r['job']))
    job_rows = defaultdict(list)
    for r in rows:
        job_rows[norm(r['job'])].append(r)
    for group in job_rows.values():
        for name, start, end in [('เตรียมเอกสาร', 'received_first', 'complete'), ('เอกสาร → ศุลกากร', 'complete', 'customs'), ('ศุลกากร → ส่งมอบ', 'customs', 'delivered')]:
            if all(r.get(start) and r.get(end) for r in group):
                days = (max(datetime.fromisoformat(r[end]) for r in group) - min(datetime.fromisoformat(r[start]) for r in group)).total_seconds() / 86400
                if days >= 0:
                    lead[name].append(days)
    denominator = counts['on_time'] + counts['late']
    counts['on_time_rate'] = round(counts['on_time'] / denominator * 100, 1) if denominator else None
    counts['issue_jobs'] = len(issues)
    return {'counts': counts, 'carriers': [{'label': k, 'value': len(v)} for k, v in sorted(carriers.items(), key=lambda x: -len(x[1]))], 'trend': [{'label': k, 'value': len(v)} for k, v in sorted(trends.items())], 'lead': [{'label': k, 'value': round(sum(v) / len(v), 1), 'samples': len(v)} for k, v in lead.items()], 'attention': sorted(attention, key=lambda r: r['plan'])[:5], 'updated_at': last, 'as_of': today}


@app.get('/api/history')
def history():
    with db() as conn:
        rows = conn.execute('SELECT id,filename,created_at,added,updated,skipped,invalid FROM batches ORDER BY created_at DESC, rowid DESC LIMIT 100').fetchall()
    return {'rows': [dict(r) for r in rows]}


@app.get('/api/history/{batch_id}')
def revisions(batch_id: str):
    with db() as conn:
        rows = conn.execute('SELECT old_data,new_data,created_at FROM revisions WHERE batch_id = ? ORDER BY id LIMIT 100', [batch_id]).fetchall()
    return {'rows': [{'before': json.loads(r['old_data']) if r['old_data'] else None, 'after': json.loads(r['new_data']), 'at': r['created_at']} for r in rows]}


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')):
        return "'" + value
    return value if value is not None else ''


@app.get('/api/export/{kind}')
def export(kind: str, request: Request):
    if kind not in ('csv', 'xlsx'):
        raise HTTPException(404, 'รูปแบบไม่รองรับ')
    with db() as conn:
        rows = selected(conn, request.query_params)
        fields = all_fields(conn)
    # Keep the source receipt column as well as the two parsed dates for round-trip imports.
    fields = [k for k, _ in FIELDS if k != 'declaration_status'] + [k for k in fields if k not in dict(FIELDS)]
    headers = [safe_cell(label(k)) for k in fields]
    data = [[safe_cell(r.get(k, '')) for k in fields] for r in rows]
    if kind == 'csv':
        buffer = io.StringIO(newline='')
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(data)
        content = buffer.getvalue().encode('utf-8-sig')
        mime = 'text/csv; charset=utf-8'
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = 'Filtered records'
        ws.append(headers)
        for row in data:
            ws.append(row)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='12695D')
        from openpyxl.utils import get_column_letter
        for i in range(1, len(fields) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 23
        totals = wb.create_sheet('Job totals')
        totals.append(['Job I', 'Cartons (once per Job)', 'CBM (once per Job)'])
        jobs = {norm(r['job']): r for r in rows}
        for r in jobs.values():
            totals.append([safe_cell(r['job']), r.get('cartons'), r.get('cbm')])
        notes = wb.create_sheet('Read me')
        notes.append(['Cartons and CBM are job totals repeated on detail rows. Sum the Job totals sheet, not detail rows.'])
        notes.append(['Job totals include the full total of each Job represented by the current filter.'])
        notes.append(['Exported at (Asia/Bangkok)', now()])
        out = io.BytesIO()
        wb.save(out)
        content = out.getvalue()
        mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return Response(content, media_type=mime, headers={'Content-Disposition': f'attachment; filename="harbor-report-{datetime.now(TZ):%Y%m%d}.{kind}"'})


app.mount('/', StaticFiles(directory=Path(__file__).parent / 'static', html=True), name='static')
